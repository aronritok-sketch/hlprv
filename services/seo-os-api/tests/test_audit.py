"""5. ütem: Screaming Frog exportok (kézi feltöltés és ügynök), technikai audit, helyszíni ellenőrzés, fejlesztői feladatok,
audit-dokumentum, DataForSEO / Ahrefs API (helyettesítő HTTP-vel), integrációk tesztje.

A tesztfájlok szintetikusak, de pontosan a Screaming Frog export formátumát követik (fejlécek, „1 - Fül - Szűrő” munkalapnév,
Inlinks szerkezet a „képek 100 KB felett” bulk exportnál)."""

import io
import json
import zipfile
from urllib.parse import quote

import httpx
import pytest
from conftest import AGENT_TOKEN, SignedClient, make_project
from openpyxl import Workbook, load_workbook

D = "https://www.example-kitchens.com"


def xlsx_bytes(sheet: str, rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(rows: list[list]) -> bytes:
    import csv

    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def internal_all(extra_ok=True, broken=True):
    h = ["Address", "Content Type", "Status Code", "Status", "Indexability", "Indexability Status", "Title 1", "Meta Description 1", "H1-1",
         "Canonical Link Element 1", "Redirect URL"]
    rows = [h,
            [f"{D}/", "text/html; charset=UTF-8", 200, "OK", "Indexable", "", "Kitchen Remodeling Naples | Example", "Desc", "Kitchen Remodeling", f"{D}/", ""],
            [f"{D}/Cabinets_Old/", "text/html; charset=UTF-8", 200, "OK", "Indexable", "", "Cabinets | Example", "", "Cabinets", "", ""],
            [f"{D}/cabinets/", "text/html; charset=UTF-8", 200, "OK", "Indexable", "", "Cabinets | Example", "Desc 2", "", f"{D}/cabinets/", ""],
            [f"{D}/promo", "text/html", 302, "Found", "Non-Indexable", "Redirected", "", "", "", "", f"{D}/"],
            [f"{D}/logo.png", "image/png", 200, "OK", "Indexable", "", "", "", "", "", ""]]
    if broken:
        rows.append([f"{D}/old-page/", "text/html", 404, "Not Found", "Non-Indexable", "Client Error", "", "", "", "", ""])
    return rows


def sf_zip(broken=True) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("internal_all.csv", csv_bytes(internal_all(broken=broken)))
        z.writestr("page_titles_duplicate.xlsx", xlsx_bytes("1 - Page Titles - Duplicate", [
            ["Address", "Occurrences", "Title 1", "Title 1 Length", "Title 1 Pixel Width", "Indexability", "Indexability Status"],
            [f"{D}/Cabinets_Old/", 1, "Cabinets | Example", 18, 160, "Indexable", None],
            [f"{D}/cabinets/", 1, "Cabinets | Example", 18, 160, "Indexable", None]]))
        z.writestr("images over 100kb.xlsx", xlsx_bytes("1 - Inlinks", [
            ["Type", "From", "To", "Anchor Text", "Alt Text", "Follow", "Target", "Rel", "Status Code", "Status"],
            ["Image", f"{D}/", f"{D}/img/hero.jpg", None, None, True, None, None, 200, "OK"],
            ["Image", f"{D}/cabinets/", f"{D}/img/hero.jpg", None, None, True, None, None, 200, "OK"],
            ["Image", f"{D}/cabinets/", f"{D}/img/cab.jpg", None, "cabinet", True, None, None, 200, "OK"]]))
        z.writestr("readme.txt", "nem export")
    return buf.getvalue()


def upload(c, pid, name, data, run_id=None):
    path = f"/projects/{pid}/crawls/upload" + (f"?run_id={run_id}" if run_id else "")
    r = c.request("POST", path, raw=data, headers={"X-Filename": quote(name), "Content-Type": "application/octet-stream"})
    assert r.status_code == 201, r.text
    return r.json()


def by_key(c, pid):
    return {f["issue_key"]: f for t in c.get(f"/projects/{pid}/audit").json()["topics"] for f in t["findings"]}


def test_manual_upload_import_and_rerun(client, as_role):
    pid = make_project(client, domain="example-kitchens.com")["id"]
    run = upload(client, pid, "sf-export.zip", sf_zip())
    rec = {f["name"]: f for f in run["files"][0]["recognized"]}
    assert rec["page_titles_duplicate.xlsx"]["issue_key"] == "title_duplicate"
    assert rec["images over 100kb.xlsx"]["issue_key"] == "images_large"
    assert rec["internal_all.csv"]["label"].startswith("Internal")
    # Második fájl ugyanahhoz a futáshoz
    run = upload(client, pid, "response_codes_client_error_(4xx).csv", csv_bytes([
        ["Address", "Content Type", "Status Code", "Status", "Indexability", "Indexability Status", "Inlinks"],
        [f"{D}/old-page/", "text/html", 404, "Not Found", "Non-Indexable", "Client Error", 3]]), run_id=run["id"])
    assert len(run["files"]) == 2
    job = client.post(f"/crawls/{run['id']}/import").json()
    assert job["status"] == "done", job
    f = by_key(client, pid)
    assert f["status_4xx"]["count"] == 1 and "Inlinks: 3" in f["status_4xx"]["sample"][0]["detail"]
    assert f["status_302"]["count"] == 1
    assert f["title_duplicate"]["count"] == 2 and f["title_duplicate"]["label"] == "2 oldal nem egyedi oldalcímmel rendelkezik"
    assert f["images_large"]["count"] == 2  # egyedi kép-URL-ek; a forrásoldal a részletben
    assert f["h1_missing"]["count"] == 1 and f["meta_missing"]["count"] == 1 and f["canonical_missing"]["count"] == 1
    assert f["url_uppercase"]["count"] == 1 and f["url_underscores"]["count"] == 1
    ov = client.get(f"/projects/{pid}/audit").json()
    assert ov["last_run"]["stats"]["html"] == 5 and ov["last_run"]["stats"]["indexable"] == 3
    sizes = {t["topic"]: t["size"] for t in ov["topics"]}
    assert sizes["http_status"] == "XL" and sizes["canonicals"] == "M" and sizes["meta"] == "S"

    # Újrafuttatás javítás után: a 404 eltűnt → javítva, előző darabszám megmarad
    run2 = upload(client, pid, "sf-export-2.zip", sf_zip(broken=False))
    assert client.post(f"/crawls/{run2['id']}/import").json()["status"] == "done"
    f = by_key(client, pid)
    assert f["status_4xx"]["count"] == 0 and f["status_4xx"]["status"] == "fixed" and f["status_4xx"]["previous_count"] == 1

    # Ismeretlen fájl: érthető hiba
    bad = upload(client, pid, "random.csv", csv_bytes([["Foo", "Bar"], [1, 2]]))
    j = client.post(f"/crawls/{bad['id']}/import").json()
    assert j["status"] == "failed" and "Screaming Frog" in j["error"]


def test_topics_findings_tasks_export_and_document(client, as_role):
    pid = make_project(client, domain="example-kitchens.com")["id"]
    run = upload(client, pid, "sf.zip", sf_zip())
    client.post(f"/crawls/{run['id']}/import")
    t = client.patch(f"/projects/{pid}/audit/topics/meta", json={"size": "M", "observation": "A blogcikkeknél hiányzik a leírás."}).json()
    assert t["size"] == "M"
    assert client.patch(f"/projects/{pid}/audit/topics/speed", json={"metrics": {"mobile": 28, "desktop": 71, "cwv_pass": False}}).status_code == 200
    assert client.patch(f"/projects/{pid}/audit/topics/meta", json={"size": "XXL"}).status_code == 422
    f = by_key(client, pid)
    dev = as_role("developer")
    assert dev.patch(f"/audit/findings/{f['status_4xx']['id']}", json={"status": "fixed"}).json()["status"] == "fixed"
    assert dev.patch(f"/audit/findings/{f['status_4xx']['id']}", json={"notes": "x"}).status_code == 403
    assert as_role("content_manager").get(f"/projects/{pid}/audit").status_code == 403
    client.patch(f"/audit/findings/{f['status_4xx']['id']}", json={"status": "open"})
    full = client.get(f"/audit/findings/{f['title_duplicate']['id']}").json()
    assert len(full["urls"]) == 2

    res = client.post(f"/projects/{pid}/audit/tasks").json()
    assert res["created"] >= 5
    tasks = [t for t in client.get(f"/projects/{pid}/tasks?role=developer").json() if t["title"].startswith("[Audit]")]
    t404 = next(t for t in tasks if "4xx" in t["title"])
    assert t404["priority"] == "XL" and f"{D}/old-page/" in t404["action"] and "4xx/5xx" in t404["done_when"]
    assert client.post(f"/projects/{pid}/audit/tasks").json()["created"] == 0

    wb = load_workbook(io.BytesIO(client.get(f"/projects/{pid}/audit/export.xlsx").content))
    assert wb.sheetnames[0] == "Összefoglaló" and "status_4xx" in wb.sheetnames

    doc = client.post(f"/projects/{pid}/documents/generate", json={"doc_type": "tech_audit"}).json()
    assert doc["status"] == "done", doc
    d = client.get(f"/documents/{doc['result']['document_id']}").json()
    assert d["language"] == "hu"
    titles = [s["title"] for s in d["content"]["sections"]]
    assert titles[0] == "Prioritási lista" and "HTTP státuszkódok (XL)" in titles and "Meta leírások (M)" in titles
    txt = json.dumps(d["content"], ensure_ascii=False)
    assert "Mit látunk a(z) example-kitchens.com esetében?" in txt and "mobil pontszáma 28/100" in txt
    assert "narrative_slot" not in txt
    # A fejlesztő láthatja az audit-dokumentumot
    assert dev.get(f"/documents/{d['id']}").status_code == 200


def agent_client():
    c = SignedClient("agent", wp_user=0)
    c.extra_headers["Authorization"] = "Bearer " + AGENT_TOKEN
    return c


def test_agent_flow(client):
    pid = make_project(client, domain="example-kitchens.com")["id"]
    agent = agent_client()
    bad = SignedClient("agent", wp_user=0)
    bad.extra_headers["Authorization"] = "Bearer rossz"
    assert bad.post("/agent/heartbeat", json={"agent": "office-pc"}).status_code == 401
    assert agent.post("/agent/heartbeat", json={"agent": "office-pc", "info": {"sf_version": "22.0", "os": "Windows"}}).json()["queued"] is False
    # Nincs feladat
    assert agent.post("/agent/crawls/claim", json={"agent": "office-pc"}).json()["run"] is None
    run = client.post(f"/projects/{pid}/crawls", json={}).json()
    assert run["status"] == "queued" and run["start_url"] == "https://example-kitchens.com/"
    assert run["agents"][0]["online"] is True
    assert client.post(f"/projects/{pid}/crawls", json={}).status_code == 409
    claimed = agent.post("/agent/crawls/claim", json={"agent": "office-pc"}).json()["run"]
    assert claimed["id"] == run["id"] and "Internal:All" in claimed["export_tabs"] and claimed["save_crawl"] is True
    assert agent.post(f"/agent/crawls/{run['id']}/progress", json={"message": "Crawl: 120 URL"}).json()["cancelled"] is False
    assert client.get(f"/projects/{pid}/crawls").json()["runs"][0]["message"] == "Crawl: 120 URL"
    up = agent.request("POST", f"/agent/crawls/{run['id']}/upload", raw=sf_zip(), headers={"X-Filename": "crawl.zip", "Content-Type": "application/zip"})
    assert up.status_code == 200 and up.json()["job"]["status"] == "done", up.text
    runs = client.get(f"/projects/{pid}/crawls").json()["runs"]
    assert runs[0]["status"] == "done" and runs[0]["agent"] == "office-pc"
    assert by_key(client, pid)["status_4xx"]["count"] == 1
    # Hiba jelentése
    run2 = client.post(f"/projects/{pid}/crawls", json={"start_url": "https://example-kitchens.com/blog/"}).json()
    agent.post("/agent/crawls/claim", json={"agent": "office-pc"})
    agent.post(f"/agent/crawls/{run2['id']}/fail", json={"error": "Licence expired"})
    assert client.get(f"/projects/{pid}/crawls").json()["runs"][0]["error"] == "Licence expired"
    # Felhasználó nem hívhatja az ügynök végpontjait
    assert client.post("/agent/crawls/claim", json={"agent": "x"}).status_code == 401


def test_site_checks(client):
    from app.services import audit as audit_service

    pid = make_project(client, domain="example-kitchens.com")["id"]

    def handler(req: httpx.Request):
        u = str(req.url)
        if u == "http://example-kitchens.com/":
            return httpx.Response(301, headers={"location": "https://example-kitchens.com/"})
        if u.endswith("/robots.txt"):
            return httpx.Response(200, text="User-agent: *\nSitemap: https://example-kitchens.com/sitemap_index.xml\n")
        if u.endswith("/sitemap_index.xml"):
            return httpx.Response(200, text="<?xml version='1.0'?><sitemapindex/>")
        return httpx.Response(404)

    audit_service._transport = httpx.MockTransport(handler)
    try:
        res = client.post(f"/projects/{pid}/audit/site-checks").json()
    finally:
        audit_service._transport = None
    assert res == {"https_redirect": True, "security_txt": False, "robots_txt": True, "sitemap": True, "sitemap_url": "https://example-kitchens.com/sitemap_index.xml"}
    f = by_key(client, pid)
    assert f["security_txt_missing"]["label"].startswith("Nem találtunk security.txt") and "sitemap_missing" not in f


# ── DataForSEO / Ahrefs ──────────────────────────────────


def dfs_ok(result, cost=0.05):
    return {"status_code": 20000, "status_message": "Ok.", "cost": cost, "tasks": [{"status_code": 20000, "status_message": "Ok.", "result": result}]}


@pytest.fixture
def api_mock(client):
    from app.services import providers

    calls = []

    def handler(req: httpx.Request):
        path = req.url.path
        calls.append(path)
        body = json.loads(req.content or b"[]") if req.method == "POST" else None
        if path == "/v3/keywords_data/google_ads/search_volume/live":
            kws = body[0]["keywords"]
            return httpx.Response(200, json=dfs_ok([{"keyword": k, "search_volume": 1000 + i, "cpc": 4.5, "competition": "HIGH",
                                                      "monthly_searches": [{"year": 2026, "month": m, "search_volume": 900 + m} for m in range(12, 0, -1)]}
                                                     for i, k in enumerate(kws)]))
        if path == "/v3/dataforseo_labs/google/bulk_keyword_difficulty/live":
            return httpx.Response(200, json=dfs_ok([{"items": [{"keyword": k, "keyword_difficulty": 33} for k in body[0]["keywords"]]}]))
        if path == "/v3/dataforseo_labs/google/keyword_suggestions/live":
            seed = body[0]["keyword"]
            return httpx.Response(200, json=dfs_ok([{"items": [
                {"keyword": seed + " cost", "keyword_info": {"search_volume": 5400, "cpc": 3.1, "competition_level": "MEDIUM"},
                 "keyword_properties": {"keyword_difficulty": 21}, "search_intent_info": {"main_intent": "informational"}},
                {"keyword": "diy " + seed, "keyword_info": {"search_volume": 900}, "keyword_properties": {}}]}]))
        if path == "/v3/serp/google/organic/live/advanced":
            return httpx.Response(200, json=dfs_ok([{"items": [
                {"type": "organic", "rank_group": 1, "domain": "www.competitor.com", "url": "https://www.competitor.com/a"},
                {"type": "paid", "rank_group": 1, "domain": "ads.com"},
                {"type": "organic", "rank_group": 2, "domain": "imperialkitchens.com", "url": "https://imperialkitchens.com/"}]}]))
        if path == "/v3/appendix/user_data":
            return httpx.Response(200, json=dfs_ok([{"login": "hpv", "money": {"balance": 42.5}}], cost=0))
        if path == "/v3/site-explorer/organic-keywords":
            return httpx.Response(200, json={"keywords": [{"keyword": "kitchen remodeling naples", "volume": 880, "keyword_difficulty": 12, "cpc": 1850,
                                                           "best_position": 4, "best_position_url": "https://imperialkitchens.com/naples/", "sum_traffic": 55}]})
        if path == "/v3/subscription-info/limits-and-usage":
            return httpx.Response(200, json={"limits_and_usage": {}})
        return httpx.Response(404, json={"error": "not mocked"})

    client.put("/settings", json={"dataforseo_login": "hpv", "dataforseo_password": "secret", "ahrefs_api_key": "ahrefs-key"})
    providers._transport = httpx.MockTransport(handler)
    yield calls
    providers._transport = None


def test_dataforseo_and_ahrefs(client, api_mock, as_role):
    p = make_project(client, domain="imperialkitchens.com", seed_keywords=[{"keyword": "kitchen remodel", "kind": "service"}])
    pid = p["id"]
    st = client.get(f"/projects/{pid}/research/api").json()
    assert st["dataforseo"] and st["ahrefs"] and st["defaults"]["location_name"] == "United States"
    # Ötletek a seed kulcsszóból – a kizárási szabály (DIY) él
    client.patch(f"/projects/{pid}", json={"excluded_topics": ["diy"]})
    res = client.post(f"/projects/{pid}/research/api", json={"action": "ideas"}).json()
    assert res["status"] == "done", res
    kws = {k["term"]: k for k in client.get(f"/projects/{pid}/keywords?include_excluded=true").json()["keywords"]}
    assert kws["kitchen remodel cost"]["volume"] == 5400 and kws["kitchen remodel cost"]["kd"] == 21
    assert kws["diy kitchen remodel"]["is_excluded"] is True
    # Volumen + nehézség az összes kulcsszóra, egy lokációhoz rendelve
    res = client.post(f"/projects/{pid}/research/api", json={"action": "volume", "location_name": "Naples,Florida,United States", "location_label": "Naples"}).json()
    assert res["status"] == "done" and res["result"]["updated"] >= 1
    # Ugyanaz a kérés gyorsítótárból: nincs új hívás
    n = len(api_mock)
    client.post(f"/projects/{pid}/research/api", json={"action": "volume", "location_name": "Naples,Florida,United States", "location_label": "Naples"})
    assert len(api_mock) == n
    # SERP: versenytársak rangsora
    kid = kws["kitchen remodel cost"]["id"]
    res = client.post(f"/projects/{pid}/research/api", json={"action": "serp", "keyword_ids": [kid]}).json()
    assert res["result"]["rankings"] == 2
    # Ahrefs organikus kulcsszavak (a CPC centben jön)
    res = client.post(f"/projects/{pid}/research/api", json={"action": "ahrefs_organic"}).json()
    assert res["status"] == "done", res
    kws = {k["term"]: k for k in client.get(f"/projects/{pid}/keywords").json()["keywords"]}
    assert kws["kitchen remodeling naples"]["own_position"] == 4
    # Költségnapló és keret
    assert client.get(f"/projects/{pid}/research/api").json()["spend_month_usd"] > 0
    client.put("/settings", json={"monthly_budget_usd": 0.01})
    res = client.post(f"/projects/{pid}/research/api", json={"action": "ideas", "seeds": ["cabinet refacing"]}).json()
    assert res["status"] == "failed" and "kerete" in res["error"]
    # Jogosultság: content manager nem futtathat API-lekérést
    assert as_role("content_manager").post(f"/projects/{pid}/research/api", json={"action": "ideas"}).status_code == 403
    # Integrációk tesztje
    t = client.post("/integrations/test").json()
    assert t["dataforseo"]["ok"] and "42.5" in t["dataforseo"]["message"]
    assert t["ahrefs"]["ok"] and t["openai"]["configured"] is False
    assert t["screaming_frog"]["ok"] is False
