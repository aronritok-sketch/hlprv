"""2. ütem: import-felismerés a valódi exportokon, összevonás, kizárás, kulcsszólista, XLSX, stílusbeli irányelvek."""

import io
from urllib.parse import quote

import openpyxl
from conftest import FIXTURES, make_project

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(c, pid, name, data=None, ctype=XLSX):
    data = data if data is not None else (FIXTURES / name).read_bytes()
    r = c.request("POST", f"/projects/{pid}/imports", raw=data, headers={"X-Filename": quote(name), "Content-Type": ctype})
    assert r.status_code == 201, r.text
    return r.json()


def hpv_project(c, **over):
    base = dict(
        name="HelloProVision",
        domain="helloprovision.com",
        locations=["Fort Myers", "Cape Coral", "Naples", "Bonita Springs", "Clearwater", "Tampa"],
        scope=["keyword_research", "structure", "content_strategy"],
        seed_keywords=[{"keyword": "logo design", "kind": "excluded"}],
        competitors=[{"domain": "theorganicmediagroup.com"}, {"domain": "prioritymarketing.com"}],
        client={"name": "HelloProVision"},
    )
    base.update(over)
    return make_project(c, **base)


def test_detect_keyword_planner(client):
    p = hpv_project(client)
    r = upload(client, p["id"], "gkp_naples.xlsx")
    d = r["detection"]
    assert d["source"] == "gkp"
    assert d["header_row"] == 2  # két címsor a fejléc előtt
    assert d["column_map"]["term"] == "Keyword"
    assert d["column_map"]["volume"] == "Avg. monthly searches"
    assert d["location"] == "Naples"  # a fájlnévből
    assert d["row_count"] > 10


def test_detect_ahrefs_exports(client):
    p = hpv_project(client)
    mt = upload(client, p["id"], "ahrefs_matching_terms.xlsx")["detection"]
    assert mt["source"] == "ahrefs_matching_terms"
    assert {"term", "volume", "kd", "parent_topic", "intents", "traffic_potential", "serp_features"} <= set(mt["column_map"])
    gap = upload(client, p["id"], "ahrefs_content_gap.xlsx")["detection"]
    assert gap["source"] == "ahrefs_content_gap"
    assert len(gap["competitor_map"]) == 10
    assert gap["competitor_map"]["exploritech.com"] == {"url": "www.exploritech.com/: URL", "position": "www.exploritech.com/: Organic Position"}


def test_detect_house_research(client):
    p = hpv_project(client)
    d = upload(client, p["id"], "house_keyword_research.xlsx")["detection"]
    assert d["source"] == "house_research"
    assert d["sheet"].startswith("Kulcsszókutatás")
    assert d["column_map"]["location"] == "Location"
    assert d["column_map"]["modifier"] == "Local long tail keywords"
    # „imagebuildingmedia.com/ Position” – kettőspont nélküli fejléc is felismerve
    assert set(d["competitor_map"]["imagebuildingmedia.com"]) == {"position", "url", "traffic"}


def run_import(c, pid, name):
    up = upload(c, pid, name)
    r = c.post(f"/projects/{pid}/imports/{up['import']['id']}/run")
    assert r.status_code == 200, r.text
    assert r.json()["job"]["status"] == "done", r.json()
    return r.json()


def test_import_merge_and_locations(client):
    p = hpv_project(client)
    pid = p["id"]
    a = run_import(client, pid, "gkp_naples.xlsx")
    assert a["import"]["status"] == "done"
    assert a["import"]["stats"]["created"] > 10
    run_import(client, pid, "gkp_fort_myers.xlsx")
    data = client.get(f"/projects/{pid}/keywords").json()
    kws = {k["term"].lower(): k for k in data["keywords"]}
    both = [k for k in kws.values() if set(k["locations"]) >= {"Naples", "Fort Myers"}]
    assert both, "ugyanaz a kulcsszó két városi exportból összevonva"
    assert "Naples" in data["locations"] and "Fort Myers" in data["locations"]
    k = both[0]
    assert {m["location"] for m in k["metrics"]} >= {"Naples", "Fort Myers"}
    # A Keyword Planner havi keresési oszlopai trendként
    assert client.get(f"/projects/{pid}/research/summary").json()["imports"] == 2


def test_import_house_research_with_competitors(client):
    p = hpv_project(client)
    pid = p["id"]
    run_import(client, pid, "house_keyword_research.xlsx")
    data = client.get(f"/projects/{pid}/keywords").json()
    kws = {k["term"]: k for k in data["keywords"]}
    k = kws["seo agency near me"]
    assert k["modifier"] == "near me"
    assert k["parent_topic"] == "seo"
    assert {"Fort Myers", "Naples", "Tampa"} <= set(k["locations"])
    fm = next(m for m in k["metrics"] if m["location"] == "Fort Myers")
    assert fm["volume"] == 8400 and fm["kd"] == 0
    assert any(c["domain"] == "theorganicmediagroup.com" and c["position"] == 20 for c in k["competitors"])
    # A „ClearWater” lokáció a projekt „Clearwater” nevéhez igazodik
    assert "Clearwater" in data["locations"]
    # Kizárt téma (logo design) automatikusan kizárva, indoklással
    logo = kws["custom logo design"]
    assert logo["is_excluded"] and logo["exclusion_reason"] == "Kizárt téma: logo design"
    assert kws["cape coral website design"]["term_location"] == "Cape Coral"


def test_import_ahrefs_intents_and_trend(client):
    p = hpv_project(client)
    run_import(client, p["id"], "ahrefs_matching_terms.xlsx")
    kws = {k["term"]: k for k in client.get(f"/projects/{p['id']}/keywords").json()["keywords"]}
    k = kws["tampa seo company"]
    assert k["source_intents"] == ["Informational", "Commercial", "Non-branded", "Local"]
    assert k["volume"] == 900 and k["kd"] == 4 and k["traffic_potential"] == 1500
    assert k["term_location"] == "Tampa"


def test_utf16_tsv_ahrefs_csv(client):
    p = hpv_project(client)
    tsv = "Keyword\tVolume\tKD\tCPC\tParent Keyword\nkitchen remodeling naples\t1,200\t12\t$4.50\tkitchen remodeling\n"
    up = upload(client, p["id"], "export.csv", tsv.encode("utf-16"), "text/csv")
    assert up["detection"]["source"] == "ahrefs_keywords"
    client.post(f"/projects/{p['id']}/imports/{up['import']['id']}/run")
    k = client.get(f"/projects/{p['id']}/keywords").json()["keywords"][0]
    assert (k["volume"], k["kd"], k["cpc"], k["term_location"]) == (1200, 12, 4.5, "Naples")


def test_screaming_frog_routed_to_audit(client):
    p = hpv_project(client)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 - H1 - Duplicate"
    ws.append(["Address", "Occurrences", "H1-1", "H1-1 Length", "Indexability"])
    ws.append(["https://example.com/a", 1, "Same", 4, "Indexable"])
    buf = io.BytesIO()
    wb.save(buf)
    up = upload(client, p["id"], "h1_duplicate.xlsx", buf.getvalue())
    assert up["detection"]["source"] == "screaming_frog"
    assert up["detection"]["sf_issue"] == "H1:Duplicate"
    assert client.post(f"/projects/{p['id']}/imports/{up['import']['id']}/run").status_code == 409


def test_bulk_actions_and_permissions(client, as_role):
    p = hpv_project(client)
    pid = p["id"]
    run_import(client, pid, "gkp_naples.xlsx")
    ids = [k["id"] for k in client.get(f"/projects/{pid}/keywords").json()["keywords"]][:3]
    cm = as_role("content_manager")
    assert cm.get(f"/projects/{pid}/keywords").status_code == 200
    assert cm.post(f"/projects/{pid}/keywords/bulk", json={"ids": ids, "action": "exclude"}).status_code == 403
    assert as_role("designer").get(f"/projects/{pid}/keywords").status_code == 403
    r = client.post(f"/projects/{pid}/keywords/bulk", json={"ids": ids, "action": "exclude", "value": "Nem releváns"})
    assert r.json()["changed"] == 3
    rows = {k["id"]: k for k in client.get(f"/projects/{pid}/keywords").json()["keywords"]}
    assert all(rows[i]["is_excluded"] and rows[i]["exclusion_reason"] == "Nem releváns" for i in ids)
    client.post(f"/projects/{pid}/keywords/bulk", json={"ids": ids[:1], "action": "include"})
    r = client.patch(f"/projects/{pid}/keywords/{ids[1]}", json={"translation": "fordítás", "is_excluded": False})
    assert r.json()["translation"] == "fordítás" and not r.json()["is_excluded"]
    assert client.post(f"/projects/{pid}/keywords/bulk", json={"ids": ids[2:], "action": "delete"}).json()["changed"] == 1
    r = client.post(f"/projects/{pid}/keywords", json={"terms": ["kitchen remodeling naples", "Kitchen  Remodeling Naples"]})
    assert r.json()["created"] == 1


def test_keyword_research_xlsx(client):
    p = hpv_project(client)
    pid = p["id"]
    run_import(client, pid, "house_keyword_research.xlsx")
    r = client.get(f"/projects/{pid}/export/keyword-research.xlsx?translation=true")
    assert r.status_code == 200
    assert "Kulcssz%C3%B3kutat%C3%A1s" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0].startswith("Kulcsszókutatás") and wb.sheetnames[1] == "kulcsszavak lokáció nélkül"
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[1]]
    assert headers[:7] == ["Keyword", "Magyar", "Local long tail keywords", "Parent topic", "Volume", "Keyword Difficulty", "Location"]
    assert "helloprovision.com/: Position" in headers and "theorganicmediagroup.com/: Traffic" in headers
    terms = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert "custom logo design" not in terms  # kizárt kulcsszó nincs az ügyfélnek szóló táblában
    vols = [row[4].value or 0 for row in ws.iter_rows(min_row=2)]
    assert vols == sorted(vols, reverse=True)


def test_style_guide(client, as_role):
    p = hpv_project(client, scope=["content_mgmt"])
    pid = p["id"]
    g = client.get(f"/projects/{pid}/style-guide").json()
    assert len(g["questions"]) == 11 and g["brand_name"] == "HelloProVision"
    r = as_role("content_manager").put(f"/projects/{pid}/style-guide", json={"answers": {"tone": "Közvetlen", "x": "y"}, "mark_received": True})
    assert r.status_code == 200 and r.json()["answers"] == {"tone": "Közvetlen"}
    intake = {i["key"]: i["status"] for i in client.get(f"/projects/{pid}").json()["intake_items"]}
    assert intake["style_guide"] == "received"
    d = client.get(f"/projects/{pid}/style-guide.docx")
    assert d.status_code == 200 and d.content[:2] == b"PK"
    from docx import Document

    text = "\n".join(par.text for par in Document(io.BytesIO(d.content)).paragraphs)
    assert "Stílusbeli irányelvek – helloprovision.com" in text
    assert "a(z) HelloProVision márkával" in text


def test_gate_requires_keywords(client):
    p = hpv_project(client)
    pid = p["id"]
    client.post(f"/projects/{pid}/transition", json={"to": "researching"})
    gates = client.get(f"/projects/{pid}/gates").json()
    assert gates["ai_analysis_complete"] == ["Még nincs importált kulcsszó (Kutatás fül)."]
    seo = SignedClientFor("seo_manager")
    r = seo.post(f"/projects/{pid}/transition", json={"to": "ai_analysis_complete"})
    assert r.status_code == 409 and r.json()["problems"]


def SignedClientFor(role):
    from conftest import SignedClient

    return SignedClient(role)
