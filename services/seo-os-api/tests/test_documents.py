"""4. ütem: dokumentumgenerálás (sablon- és Claude-útvonal), fájlok, elavultság, jóváhagyás, referenciadokumentumok,
gyártási feladatok és CRM-összekötés, szerepkörök szerinti láthatóság."""

import io
import json
from types import SimpleNamespace
from urllib.parse import quote

import pytest
from docx import Document as Docx
from openpyxl import load_workbook
from test_intelligence import run, seed_project


def full_project(c):
    p = seed_project(c)
    pid = p["id"]
    run(c, f"/projects/{pid}/analysis/run")
    run(c, f"/projects/{pid}/structure/generate")
    run(c, f"/projects/{pid}/roadmap/generate")
    pages = {pg["url"]: pg for pg in c.get(f"/projects/{pid}/pages").json()["pages"]}
    run(c, f"/projects/{pid}/wireframes/generate", {"page_ids": [pages["/naples-kitchen-remodeling/"]["id"]]})
    return pid


def download(c, doc, fmt):
    r = c.get(f"/documents/{doc['id']}/files/{fmt}")
    assert r.status_code == 200, r.text
    return r.content


def test_template_documents_all_types(client):
    pid = full_project(client)
    listing = client.get(f"/projects/{pid}/documents").json()
    assert {t["doc_type"] for t in listing["types"]} >= {"seo_strategy", "content_strategy", "dev_brief", "writer_brief", "designer_brief", "seo_checklist"}
    ids = {}
    for t in listing["types"]:
        if t["doc_type"] == "tech_audit":
            continue  # audit-adat kell hozzá, lásd test_audit.py
        res = run(client, f"/projects/{pid}/documents/generate", {"doc_type": t["doc_type"]})
        assert res["method"] == "template" and res["version"] == 1
        ids[t["doc_type"]] = res["document_id"]
    docs = {d["doc_type"]: d for d in client.get(f"/projects/{pid}/documents").json()["documents"]}
    assert docs["seo_strategy"]["language"] == "en"  # angol tartalmi nyelvű projekt ügyféldokumentuma angol
    assert docs["dev_brief"]["language"] == "hu"
    assert all(d["stale"] is False for d in docs.values())

    pdf = download(client, docs["seo_strategy"], "pdf")
    assert pdf.startswith(b"%PDF")
    # Fejlesztői brief DOCX: a táblában ott vannak az URL-ek és a H1-ek
    dev = Docx(io.BytesIO(download(client, docs["dev_brief"], "docx")))
    cells = {c.text for t in dev.tables for row in t.rows for c in row.cells}
    assert "/naples-kitchen-remodeling/" in cells
    assert any("Belső linkek" in p.text for p in dev.paragraphs)
    # Tartalomstratégia XLSX: a házi 6 lapos munkafüzet
    wb = load_workbook(io.BytesIO(download(client, docs["content_strategy"], "xlsx")))
    assert wb.sheetnames == ["Összefoglaló", "Kulcsszóklaszterek", "6 havi roadmap", "Mérési terv", "PPC", "Következő témák"]
    assert wb["6 havi roadmap"]["A1"].value == "Month" and wb["6 havi roadmap"].max_row > 10
    # A közvetlen export is működik
    r = client.get(f"/projects/{pid}/exports/content-strategy.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"

    # Elavulás: ha a struktúra változik, a brief elavult
    page = next(pg for pg in client.get(f"/projects/{pid}/pages").json()["pages"] if pg["url"] == "/naples-kitchen-remodeling/")
    assert client.patch(f"/projects/{pid}/pages/{page['id']}", json={"h1": "Kitchen Remodeling Naples – új"}).status_code == 200
    docs = {d["doc_type"]: d for d in client.get(f"/projects/{pid}/documents").json()["documents"]}
    assert docs["dev_brief"]["stale"] is True
    # Új változat
    res = run(client, f"/projects/{pid}/documents/generate", {"doc_type": "dev_brief", "formats": ["pdf"]})
    assert res["version"] == 2
    latest = client.get(f"/documents/{res['document_id']}").json()
    assert [f["format"] for f in latest["files"]] == ["pdf"]
    assert any("Kitchen Remodeling Naples – új" in json.dumps(latest["content"], ensure_ascii=False) for _ in [0])


def test_requirements_and_workflow(client, as_role):
    p = seed_project(client)
    pid = p["id"]
    r = client.post(f"/projects/{pid}/documents/generate", json={"doc_type": "dev_brief"}).json()
    assert r["status"] == "failed" and "oldalstruktúrát" in r["error"]
    pid = full_project(client)
    doc_id = run(client, f"/projects/{pid}/documents/generate", {"doc_type": "seo_strategy"})["document_id"]
    # Szerkesztés: szöveg átírása → a fájlok újrakészülnek
    d = client.get(f"/documents/{doc_id}").json()
    content = d["content"]
    content["lead"] = "Kézzel átírt bevezető."
    old_file = d["files"][0]["file_id"]
    d2 = client.patch(f"/documents/{doc_id}", json={"content": content}).json()
    assert d2["content"]["lead"] == "Kézzel átírt bevezető." and d2["files"][0]["file_id"] != old_file
    # Kiküldés csak jóváhagyás után
    assert client.patch(f"/documents/{doc_id}", json={"status": "sent"}).status_code == 409
    cm = as_role("content_manager")
    assert cm.patch(f"/documents/{doc_id}", json={"status": "approved"}).status_code == 403
    assert client.patch(f"/documents/{doc_id}", json={"status": "approved"}).json()["approved_at"]
    assert client.patch(f"/documents/{doc_id}", json={"content": content}).status_code == 409
    assert client.patch(f"/documents/{doc_id}", json={"status": "sent"}).json()["sent_at"]
    # Láthatóság: a fejlesztő nem látja az ügyfél SEO stratégiát, és nem generálhat
    dev = as_role("developer")
    assert dev.get(f"/documents/{doc_id}").status_code == 403
    assert all(d["doc_type"] in ("dev_brief", "seo_checklist", "tech_audit") for d in dev.get(f"/projects/{pid}/documents").json()["documents"])
    assert dev.post(f"/projects/{pid}/documents/generate", json={"doc_type": "dev_brief"}).status_code == 403
    # A content manager a szövegírói briefet generálhatja, a fejlesztőit nem
    assert cm.post(f"/projects/{pid}/documents/generate", json={"doc_type": "writer_brief"}).json()["status"] == "done"
    assert cm.post(f"/projects/{pid}/documents/generate", json={"doc_type": "dev_brief"}).status_code == 403


class FakeStream:
    def __init__(self, message):
        self.message = message

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def get_final_message(self):
        return self.message


class FakeClaude:
    """Az Anthropic SDK helyettesítője: a kapott sémából érvényes szakaszszöveget ad vissza."""

    def __init__(self, calls, stop_reason="end_turn"):
        self.calls = calls
        self.stop_reason = stop_reason
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self._stream))

    def _stream(self, **kw):
        self.calls.append(kw)
        keys = kw["output_config"]["format"]["schema"]["properties"]["sections"]["items"]["properties"]["key"]["enum"]
        out = {"lead": "Claude bevezető.", "sections": [{"key": k, "paragraphs": [f"Claude szöveg: {k}."], "bullets": []} for k in keys]}
        msg = SimpleNamespace(
            content=[SimpleNamespace(type="thinking", thinking=""), SimpleNamespace(type="text", text=json.dumps(out))],
            stop_reason=self.stop_reason, model=kw["model"],
            usage=SimpleNamespace(input_tokens=1000, output_tokens=400, cache_read_input_tokens=0, cache_creation_input_tokens=0),
        )
        return FakeStream(msg)


@pytest.fixture
def claude_on(client):
    from app.config import get_settings
    from app.services import claude

    calls = []
    client.put("/settings", json={"anthropic_api_key": "sk-ant-test-1234"})
    get_settings().llm_fake = False
    state = {"stop": "end_turn"}
    claude._client_factory = lambda key: FakeClaude(calls, state["stop"])
    yield calls, state
    get_settings().llm_fake = True
    claude._client_factory = None


def test_claude_path_with_reference_docs(client, claude_on, db):
    calls, state = claude_on
    # Saját minta feltöltése: a rendszerpromptba kerül
    ref = Docx()
    ref.add_paragraph("HPV MINTA: Így írunk fejlesztői briefet – konkrétan, elfogadási feltétellel.")
    buf = io.BytesIO()
    ref.save(buf)
    r = client.request("POST", "/reference-docs?doc_type=dev_brief&language=hu", raw=buf.getvalue(),
                       headers={"X-Filename": quote("gulyastamas.hu SEO fejlesztői módosítások.docx"), "Content-Type": "application/octet-stream"})
    assert r.status_code == 201, r.text
    assert client.get("/reference-docs").json()[0]["title"] == "gulyastamas.hu SEO fejlesztői módosítások"

    pid = full_project(client)
    res = run(client, f"/projects/{pid}/documents/generate", {"doc_type": "dev_brief"})
    assert res["method"] == "claude"
    kw = calls[-1]
    assert kw["model"] == "claude-opus-5" and kw["thinking"] == {"type": "adaptive"} and kw["fallbacks"] == "default"
    assert "HPV MINTA" in kw["system"][0]["text"] and kw["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert "/naples-kitchen-remodeling/" in kw["messages"][0]["content"]
    d = client.get(f"/documents/{res['document_id']}").json()
    assert d["content"]["lead"] == "Claude bevezető." and d["model"] == "claude-opus-5"
    texts = json.dumps(d["content"], ensure_ascii=False)
    assert "Claude szöveg: decision_basis." in texts
    # Ugyanazokkal az adatokkal gyorsítótárból jön (nincs új hívás)
    n = len(calls)
    run(client, f"/projects/{pid}/documents/generate", {"doc_type": "dev_brief"})
    assert len(calls) == n
    from sqlalchemy import select

    from app.models import ApiLog

    assert any(lg.provider == "anthropic" and lg.tokens_in == 1000 for lg in db.scalars(select(ApiLog)).all())
    # Elutasítás esetén a sablonszöveggel készül el, figyelmeztetéssel
    state["stop"] = "refusal"
    res = run(client, f"/projects/{pid}/documents/generate", {"doc_type": "writer_brief"})
    assert res["method"] == "template" and "elutasítás" in res["warning"]


def test_production_tasks_and_crm(client, as_role):
    pid = full_project(client)
    res = client.post(f"/projects/{pid}/tasks/generate").json()
    assert res["created"] > 20
    tasks = client.get(f"/projects/{pid}/tasks").json()
    roles = {t["role"] for t in tasks}
    assert roles == {"developer", "writer", "designer", "seo"}
    create = next(t for t in tasks if t["title"] == "Oldal létrehozása: /naples-kitchen-remodeling/")
    assert "sitemap" in create["done_when"] and create["role_label"] == "Fejlesztő"
    art = next(t for t in tasks if t["role"] == "writer" and t["roadmap_item_id"])
    assert art["due_date"] and art["target_url"].startswith("/insights/")
    # Ismételt generálás: nem duplikál
    again = client.post(f"/projects/{pid}/tasks/generate").json()
    assert again["created"] == 0 and len(client.get(f"/projects/{pid}/tasks").json()) == len(tasks)
    # Fejlesztő: csak a saját szerepkörének feladatán, csak státusz
    dev = as_role("developer")
    dev.get("/me")
    assert dev.patch(f"/tasks/{create['id']}", json={"status": "in_progress", "title": "x"}).json()["title"] == create["title"]
    assert dev.patch(f"/tasks/{art['id']}", json={"status": "done"}).status_code == 403
    # Hozzárendelés és „Saját munkám”
    users = {u["role"]: u for u in client.get("/users").json()}
    client.patch(f"/tasks/{create['id']}", json={"assignee_id": users["developer"]["id"]})
    mine = dev.get("/me/tasks").json()
    assert [t["id"] for t in mine["mine"]] == [create["id"]] and mine["mine"][0]["assignee_wp_id"] == 5
    assert any(t["role"] == "developer" for t in mine["open"])
    assert dev.get("/dashboard").json()["my_tasks"][0]["id"] == create["id"]
    # CRM: át nem küldött feladatok, majd összekötés
    unpushed = client.get(f"/projects/{pid}/tasks?unpushed=1").json()
    assert len(unpushed) == len(tasks)
    assert client.post(f"/projects/{pid}/tasks/crm-links", json=[{"id": create["id"], "crm_task_id": 991}]).json()["linked"] == 1
    assert len(client.get(f"/projects/{pid}/tasks?unpushed=1").json()) == len(tasks) - 1
