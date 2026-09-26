"""Táblázatok beolvasása és felismerése: Ahrefs (Keywords Explorer, Matching terms, Organic keywords, Content gap),
Google Keyword Planner, a saját „Kulcsszókutatás” tábla, Screaming Frog export, és bármilyen tábla „Keyword” oszloppal.

A felismerés eredménye egy javaslat (forrás, fejlécsor, oszlop-hozzárendelés, versenytárs-oszlopok, lokáció),
amit a felületen jóvá lehet hagyni vagy módosítani, mielőtt az import lefut.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Optional

import openpyxl

MAX_ROWS = 50_000

# Mező → a fejlécben előforduló nevek (kisbetűs, szóköz-normalizált).
SYNONYMS: dict[str, list[str]] = {
    "term": ["keyword", "keywords", "kulcsszó", "kulcsszavak", "search term", "query", "kifejezés", "top queries"],
    "volume": ["volume", "avg. monthly searches", "search volume", "volume (us)", "keresési volumen", "havi keresés", "sv"],
    "kd": ["kd", "difficulty", "keyword difficulty", "kulcsszónehézség"],
    "cpc": ["cpc", "top of page bid (high range)"],
    "traffic_potential": ["traffic potential"],
    "parent_topic": ["parent topic", "parent keyword"],
    "intents": ["intents", "intent"],
    "serp_features": ["serp features"],
    "location": ["location", "lokáció", "city", "város"],
    "modifier": ["local long tail keywords", "local long tail", "modifier"],
    "translation": ["magyar", "fordítás", "translation"],
    "category": ["category", "kategória"],
    "competition": ["competition"],
    "currency": ["currency", "pénznem"],
    "own_position": ["current position", "position"],
    "own_url": ["current url", "url"],
    "own_traffic": ["organic traffic", "traffic"],
    "sv_trend": [],
}

COMPETITOR_RE = re.compile(r"^(?P<domain>[^\s:]+\.[^\s:]+?)/?\s*:?\s*(?P<field>organic position|position|url|organic traffic|traffic)$", re.I)
GKP_MONTH_RE = re.compile(r"^searches:\s", re.I)
SF_SHEET_RE = re.compile(r"^\d+\s*-\s*(?P<tab>.+?)\s*-\s*(?P<filter>.+)$")


def norm_header(h: Any) -> str:
    return " ".join(str(h or "").replace("\n", " ").split()).strip().lower()


def normalize_term(term: str) -> str:
    t = unicodedata.normalize("NFC", str(term or "")).strip().lower()
    t = re.sub(r"\s+", " ", t)
    return t.strip(" \"'")


def fold(s: str) -> str:
    """Összehasonlításhoz: kisbetű, ékezet és szóköz nélkül."""
    s = unicodedata.normalize("NFKD", s.lower())
    return "".join(c for c in s if c.isalnum())


def to_int(v: Any) -> Optional[int]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    # „1,2K” vagy „10K – 100K” típusú értékek
    m = re.match(r"^([\d.,]+)\s*([kKmM]?)", s)
    if not m:
        return None
    num = m.group(1)
    if num.count(",") and num.count("."):
        num = num.replace(",", "")
    elif num.count(",") == 1 and len(num.split(",")[1]) != 3:
        num = num.replace(",", ".")
    else:
        num = num.replace(",", "")
    try:
        val = float(num)
    except ValueError:
        return None
    mult = {"k": 1_000, "m": 1_000_000}.get(m.group(2).lower(), 1)
    return int(round(val * mult))


def to_float(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace("$", "").replace("€", "").replace(" ", "")
    s = s.replace(",", ".") if s.count(",") == 1 and s.count(".") == 0 else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def split_list(v: Any) -> list[str]:
    if not v:
        return []
    return [x.strip() for x in re.split(r"[,;]", str(v)) if x.strip()]


@dataclass
class Sheet:
    name: str
    rows: list[list[Any]]


def read_table(path: Path, filename: str) -> list[Sheet]:
    """XLSX vagy CSV/TSV (az Ahrefs CSV exportja UTF-16, tabulátorral tagolt)."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append(list(row))
            while rows and not any(c not in (None, "") for c in rows[-1]):
                rows.pop()
            sheets.append(Sheet(ws.title, rows))
        wb.close()
        return sheets
    if lower.endswith((".csv", ".tsv", ".txt")):
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "utf-16", "cp1250", "latin-1"):
            try:
                text = raw.decode(enc)
                if enc == "utf-16" and "\x00" in text:
                    continue
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("A fájl kódolása nem ismerhető fel.")
        sample = text[:4096]
        delim = "\t" if sample.count("\t") > sample.count(",") and sample.count("\t") > sample.count(";") else (";" if sample.count(";") > sample.count(",") else ",")
        rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)][:MAX_ROWS]
        return [Sheet(Path(filename).stem, rows)]
    raise ValueError("Csak XLSX, CSV vagy TSV fájl tölthető fel.")


@dataclass
class Detection:
    source: str
    sheet: str
    header_row: int
    columns: list[str]
    column_map: dict[str, str]
    competitor_map: dict[str, dict[str, str]] = field(default_factory=dict)
    location: str = ""
    sf_issue: str = ""
    preview: list[dict[str, Any]] = field(default_factory=list)
    row_count: int = 0

    def as_dict(self) -> dict[str, Any]:
        return {
            "source": self.source,
            "sheet": self.sheet,
            "header_row": self.header_row,
            "columns": self.columns,
            "column_map": self.column_map,
            "competitor_map": self.competitor_map,
            "location": self.location,
            "sf_issue": self.sf_issue,
            "preview": self.preview,
            "row_count": self.row_count,
        }


def find_header(rows: list[list[Any]]) -> int:
    for i, row in enumerate(rows[:20]):
        cells = [norm_header(c) for c in row]
        if any(c in SYNONYMS["term"] for c in cells) or (cells and cells[0] in ("address", "type")):
            return i
    return 0


def map_columns(headers: list[str]) -> tuple[dict[str, str], dict[str, dict[str, str]]]:
    normalized = [norm_header(h) for h in headers]
    colmap: dict[str, str] = {}
    competitors: dict[str, dict[str, str]] = {}
    for raw, h in zip(headers, normalized):
        m = COMPETITOR_RE.match(h)
        if m:
            domain = m.group("domain").lower().removeprefix("www.").rstrip("/")
            fld = m.group("field").lower()
            key = "position" if "position" in fld else ("url" if fld == "url" else "traffic")
            competitors.setdefault(domain, {})[key] = str(raw)
            continue
        if h.startswith("sv trend"):
            colmap.setdefault("sv_trend", str(raw))
            continue
        for fld, names in SYNONYMS.items():
            if h in names and fld not in colmap:
                colmap[fld] = str(raw)
                break
    # Egyszerű „Position”/„URL”/„Traffic” csak akkor a saját oldalé, ha nincs versenytárs-oszlop.
    if competitors:
        for k in ("own_position", "own_url", "own_traffic"):
            if k in colmap and norm_header(colmap[k]) in ("position", "url", "traffic"):
                colmap.pop(k)
    return colmap, competitors


def guess_source(filename: str, headers: list[str], colmap: dict[str, str], competitors: dict, sheet_name: str) -> str:
    h = {norm_header(x) for x in headers}
    fn = filename.lower()
    if "address" in h or ("type" in h and "from" in h and "to" in h) or SF_SHEET_RE.match(sheet_name or ""):
        return "screaming_frog"
    if "avg. monthly searches" in h or "keyword_planner" in fn or "keyword stats" in fn:
        return "gkp"
    if competitors and "location" in colmap:
        return "house_research"
    if competitors:
        return "ahrefs_content_gap"
    if "matching-terms" in fn or "matching_terms" in fn:
        return "ahrefs_matching_terms"
    if "current position" in h or "organic-keywords" in fn or "organic_keywords" in fn:
        return "ahrefs_organic"
    if "parent keyword" in h or "traffic potential" in h:
        return "ahrefs_keywords"
    return "manual"


def location_from_filename(filename: str, locations: list[str]) -> str:
    folded = fold(Path(filename).stem.replace("_", " "))
    for loc in sorted(locations, key=len, reverse=True):
        if loc and fold(loc) and fold(loc) in folded:
            return loc
    return ""


def match_location(value: str, locations: list[str]) -> str:
    """A táblában szereplő település hozzáigazítása a projekt lokációihoz (pl. „ClearWater” → „Clearwater”)."""
    v = fold(value or "")
    if not v:
        return ""
    for loc in locations:
        lf = fold(loc)
        if lf and (lf == v or lf.startswith(v) or v.startswith(lf)):
            return loc
    return str(value).strip()


def term_location(term: str, locations: list[str]) -> str:
    t = " " + normalize_term(term) + " "
    for loc in sorted(locations, key=len, reverse=True):
        ln = normalize_term(loc)
        if ln and f" {ln} " in t:
            return loc
    return ""


def detect(sheets: list[Sheet], filename: str, locations: list[str]) -> Detection:
    """A legtöbb adatot tartalmazó, felismerhető munkalap."""
    best: Optional[Detection] = None
    for sheet in sheets:
        if not sheet.rows:
            continue
        hr = find_header(sheet.rows)
        headers = [str(c) if c is not None else "" for c in sheet.rows[hr]]
        colmap, competitors = map_columns(headers)
        source = guess_source(filename, headers, colmap, competitors, sheet.name)
        data_rows = [r for r in sheet.rows[hr + 1 :] if any(c not in (None, "") for c in r)]
        det = Detection(
            source=source,
            sheet=sheet.name,
            header_row=hr,
            columns=headers,
            column_map=colmap,
            competitor_map=competitors,
            location=location_from_filename(filename, locations) if "location" not in colmap else "",
            row_count=len(data_rows),
        )
        if source == "screaming_frog":
            m = SF_SHEET_RE.match(sheet.name)
            det.sf_issue = f"{m.group('tab')}:{m.group('filter')}" if m else Path(filename).stem
        det.preview = [
            {headers[i]: _plain(r[i]) for i in range(min(len(headers), len(r))) if headers[i]} for r in data_rows[:8]
        ]
        usable = "term" in colmap or source == "screaming_frog"
        if usable and (best is None or det.row_count > best.row_count):
            best = det
    if best is None:
        raise ValueError("Nem találtam „Keyword” oszlopot egyik munkalapon sem.")
    return best


def _plain(v: Any) -> Any:
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


@dataclass
class Row:
    term: str
    location: str = ""
    volume: Optional[int] = None
    kd: Optional[int] = None
    cpc: Optional[float] = None
    traffic_potential: Optional[int] = None
    parent_topic: str = ""
    intents: list[str] = field(default_factory=list)
    serp_features: list[str] = field(default_factory=list)
    modifier: str = ""
    translation: str = ""
    category: str = ""
    competition: str = ""
    currency: str = ""
    trend: list[int] = field(default_factory=list)
    own: Optional[tuple[Optional[int], str, Optional[int]]] = None
    competitors: dict[str, tuple[Optional[int], str, Optional[int]]] = field(default_factory=dict)


def iter_rows(sheet: Sheet, det: Detection, locations: list[str]) -> Iterator[Row]:
    headers = [str(c) if c is not None else "" for c in sheet.rows[det.header_row]]
    idx = {h: i for i, h in enumerate(headers) if h}
    cm = det.column_map

    def cell(row, fld):
        col = cm.get(fld)
        if col is None or col not in idx:
            return None
        i = idx[col]
        return row[i] if i < len(row) else None

    month_cols = [i for i, h in enumerate(headers) if GKP_MONTH_RE.match(norm_header(h))]
    for raw in sheet.rows[det.header_row + 1 :]:
        term = cell(raw, "term")
        if term is None or str(term).strip() == "":
            continue
        term = " ".join(str(term).split())
        loc = det.location
        if "location" in cm:
            loc = match_location(str(cell(raw, "location") or ""), locations)
        trend: list[int] = []
        if month_cols:
            trend = [to_int(raw[i]) or 0 for i in month_cols if i < len(raw)]
        elif "sv_trend" in cm:
            trend = [to_int(x) or 0 for x in split_list(cell(raw, "sv_trend"))]
        r = Row(
            term=term,
            location=loc,
            volume=to_int(cell(raw, "volume")),
            kd=to_int(cell(raw, "kd")),
            cpc=to_float(cell(raw, "cpc")),
            traffic_potential=to_int(cell(raw, "traffic_potential")),
            parent_topic=str(cell(raw, "parent_topic") or "").strip(),
            intents=split_list(cell(raw, "intents")),
            serp_features=split_list(cell(raw, "serp_features")),
            modifier=str(cell(raw, "modifier") or "").strip(),
            translation=str(cell(raw, "translation") or "").strip(),
            category=str(cell(raw, "category") or "").strip(),
            competition=str(cell(raw, "competition") or "").strip(),
            currency=str(cell(raw, "currency") or "").strip().upper()[:8],
            trend=trend,
        )
        if any(k in cm for k in ("own_position", "own_url")):
            pos, url = to_int(cell(raw, "own_position")), str(cell(raw, "own_url") or "").strip()
            if pos or url:
                r.own = (pos, url, to_int(cell(raw, "own_traffic")))
        for domain, cols in det.competitor_map.items():
            pos = to_int(raw[idx[cols["position"]]]) if cols.get("position") in idx and idx[cols["position"]] < len(raw) else None
            url = str(raw[idx[cols["url"]]] or "").strip() if cols.get("url") in idx and idx[cols["url"]] < len(raw) else ""
            trf = to_int(raw[idx[cols["traffic"]]]) if cols.get("traffic") in idx and idx[cols["traffic"]] < len(raw) else None
            if pos or url:
                r.competitors[domain] = (pos, url, trf)
        yield r
