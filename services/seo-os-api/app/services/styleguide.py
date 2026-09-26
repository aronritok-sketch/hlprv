"""Stílusbeli irányelvek: az ügyfélnek kiküldhető, személyre szabott DOCX.

A folyamat szabálya: nem a sablont küldjük ki – a címben az ügyfél domainje szerepel, a márkanév pedig
a sablon kiemelt helyein át van írva.
"""

import io

from docx import Document
from docx.shared import Pt, RGBColor

from ..models import STYLE_QUESTIONS


def build_docx(domain: str, brand: str, answers: dict) -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    doc.add_heading(f"Stílusbeli irányelvek – {domain}", level=0)
    p = doc.add_paragraph()
    run = p.add_run(brand)
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph(
        "Kérjük, válaszoljatok az alábbi kérdésekre, hogy a cikkek a márkátok hangján, szakmailag pontosan készüljenek. "
        "Ahol nincs mit írni, a kérdés üresen maradhat."
    )
    for key, question, hint in STYLE_QUESTIONS:
        q = question.replace("a márkával", f"a(z) {brand} márkával").replace("amivel a márka", f"amivel a(z) {brand}")
        doc.add_heading(q, level=1)
        if hint:
            h = doc.add_paragraph(hint)
            for r in h.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x66)
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = str(answers.get(key, "") or "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
