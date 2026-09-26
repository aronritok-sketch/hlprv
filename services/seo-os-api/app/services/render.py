"""Dokumentumok kirajzolása: PDF (Jinja2 + WeasyPrint), DOCX (python-docx) és XLSX (openpyxl).

Mindhárom ugyanabból a tartalom-JSON-ból készül (lásd docfacts), így a formátumok tartalma azonos.
Az ügyfél-PDF a HelloProVision ajánlataihoz igazodik: sötét borító, kiemelőszín, táblák, lábléc oldalszámmal.
A belső briefek DOCX-ben a saját fejlesztői / szövegírói dokumentumok elrendezését követik (fejléc, címkés táblák).
"""

import io
from datetime import date
from typing import Any

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import Environment, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from . import xlsx

PDF_TEMPLATE = r"""<!doctype html>
<html lang="{{ lang }}"><head><meta charset="utf-8"><title>{{ c.title }}</title>
<style>
@page { size: A4 {{ 'landscape' if landscape else 'portrait' }}; margin: 18mm 16mm 20mm 16mm;
  @bottom-left { content: "{{ footer }}"; font: 8pt "DejaVu Sans", sans-serif; color: #8a8a80; }
  @bottom-right { content: counter(page) " / " counter(pages); font: 8pt "DejaVu Sans", sans-serif; color: #8a8a80; } }
@page cover { margin: 0; @bottom-left { content: none } @bottom-right { content: none } }
* { box-sizing: border-box; }
body { font: 9.6pt/1.5 "DejaVu Sans", "Liberation Sans", Arial, sans-serif; color: #1d1d1a; margin: 0; }
.cover { page: cover; height: {{ '210mm' if landscape else '297mm' }}; background: #121210; color: #f2f1ea; padding: 28mm 22mm; position: relative; page-break-after: always; }
.cover .brand { font-weight: 700; letter-spacing: .08em; text-transform: uppercase; font-size: 10pt; color: {{ accent }}; }
.cover .chip { display: inline-block; margin-top: 40mm; padding: 2mm 4mm; border: 1px solid {{ accent }}; color: {{ accent }}; border-radius: 99px; font-size: 8.5pt; letter-spacing: .1em; }
.cover h1 { font-size: 34pt; line-height: 1.1; margin: 8mm 0 4mm; font-weight: 800; }
.cover .sub { font-size: 12pt; color: #c9c7ba; }
.cover .bar { position: absolute; left: 22mm; right: 22mm; bottom: 26mm; border-top: 3px solid {{ accent }}; padding-top: 4mm; font-size: 9pt; color: #a9a79a; display: flex; justify-content: space-between; }
.cover .bar span + span { float: right; }
.cover .dot { position: absolute; right: 22mm; top: 26mm; width: 18mm; height: 18mm; border-radius: 50%; background: {{ accent2 }}; }
h2 { font-size: 16pt; margin: 0 0 3mm; padding-top: 2mm; color: #121210; page-break-after: avoid; }
h2 .n { color: {{ accent_dark }}; margin-right: 2mm; }
section { margin-bottom: 8mm; }
section.break { page-break-before: always; }
p { margin: 0 0 2.6mm; }
.lead { font-size: 11pt; color: #33332e; border-left: 3px solid {{ accent }}; padding-left: 4mm; margin-bottom: 6mm; }
ul { margin: 0 0 3mm 0; padding-left: 5mm; } li { margin-bottom: 1.2mm; }
table { width: 100%; border-collapse: collapse; margin: 2mm 0 4mm; font-size: 8.2pt; page-break-inside: auto; }
thead { display: table-header-group; }
th { background: #121210; color: #f2f1ea; text-align: left; padding: 2mm 2.2mm; font-weight: 600; vertical-align: bottom; }
td { padding: 1.8mm 2.2mm; border-bottom: .5pt solid #d8d6ca; vertical-align: top; word-wrap: break-word; }
tr { page-break-inside: avoid; }
tbody tr:nth-child(even) td { background: #f6f5ef; }
.kpis { display: flex; gap: 3mm; margin: 2mm 0 5mm; }
.kpi { flex: 1; background: #121210; color: #f2f1ea; border-radius: 3mm; padding: 4mm; }
.kpi b { display: block; font-size: 18pt; color: {{ accent }}; line-height: 1.1; }
.kpi span { font-size: 8pt; color: #c9c7ba; } .kpi i { display: block; font-style: normal; font-size: 7.5pt; color: #8a8a80; }
.cards { display: flex; flex-wrap: wrap; gap: 3mm; margin: 2mm 0 4mm; }
.card { width: calc(50% - 1.5mm); border: .6pt solid #d8d6ca; border-top: 2.5pt solid {{ accent }}; border-radius: 2mm; padding: 3mm 3.5mm; page-break-inside: avoid; }
.card b { display: block; margin-bottom: 1mm; }
.callout { background: #f3f7e8; border-left: 3pt solid {{ accent }}; padding: 3mm 4mm; margin: 2mm 0 4mm; page-break-inside: avoid; }
.callout b { display: block; margin-bottom: 1mm; }
.steps { counter-reset: s; margin: 2mm 0 4mm; }
.step { position: relative; padding: 0 0 3mm 11mm; page-break-inside: avoid; }
.step:before { counter-increment: s; content: counter(s); position: absolute; left: 0; top: 0; width: 7mm; height: 7mm; border-radius: 50%; background: {{ accent2 }}; color: #fff; text-align: center; line-height: 7mm; font-weight: 700; font-size: 8.5pt; }
</style></head><body>
<div class="cover">
  <div class="dot"></div>
  <div class="brand">{{ brand }}</div>
  <div class="chip">{{ c.chip }}</div>
  <h1>{{ c.title }}</h1>
  <div class="sub">{{ c.subtitle }}</div>
  <div class="bar"><span>{{ footer }}</span><span>{{ today }}{% if version %} · v{{ version }}{% endif %}</span></div>
</div>
{% if c.lead %}<p class="lead">{{ c.lead }}</p>{% endif %}
{% for s in c.sections %}
<section class="{{ 'break' if s.break else '' }}">
  <h2><span class="n">{{ '%02d' % loop.index }}</span>{{ s.title }}</h2>
  {% for b in s.blocks %}
    {% if b.type == 'paragraph' %}<p>{{ b.text }}</p>
    {% elif b.type == 'bullets' %}<ul>{% for i in b['items'] %}<li>{{ i }}</li>{% endfor %}</ul>
    {% elif b.type == 'callout' %}<div class="callout">{% if b.title %}<b>{{ b.title }}</b>{% endif %}{{ b.text }}</div>
    {% elif b.type == 'kpis' %}<div class="kpis">{% for k in b['items'] %}<div class="kpi"><b>{{ k.value }}</b><span>{{ k.label }}</span>{% if k.hint %}<i>{{ k.hint }}</i>{% endif %}</div>{% endfor %}</div>
    {% elif b.type == 'cards' %}<div class="cards">{% for k in b['items'] %}<div class="card"><b>{{ k.title }}</b>{{ k.text }}</div>{% endfor %}</div>
    {% elif b.type == 'steps' %}<div class="steps">{% for k in b['items'] %}<div class="step"><b>{{ k.title }}</b><br>{{ k.text }}</div>{% endfor %}</div>
    {% elif b.type == 'table' %}<table><thead><tr>{% for h in b.columns %}<th>{{ h }}</th>{% endfor %}</tr></thead>
      <tbody>{% for r in b.rows %}<tr>{% for v in r %}<td>{{ v }}</td>{% endfor %}</tr>{% endfor %}</tbody></table>
    {% endif %}
  {% endfor %}
</section>
{% endfor %}
</body></html>"""

_env = Environment(autoescape=select_autoescape(default=True, default_for_string=True))
_pdf_tpl = _env.from_string(PDF_TEMPLATE)


def _darken(hex_color: str, factor: float = 0.62) -> str:
    h = hex_color.lstrip("#")
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return "#5c7a1f"
    return "#%02x%02x%02x" % tuple(int(v * factor) for v in (r, g, b))


def _wide(content: dict) -> bool:
    return any(b.get("type") == "table" and len(b.get("columns", [])) >= 8 for s in content.get("sections", []) for b in s.get("blocks", []))


def html(content: dict, style: dict, version: int = 0) -> str:
    accent = style.get("accent") or "#A4DA4C"
    return _pdf_tpl.render(
        c=content, lang=style.get("lang", "hu"), brand=style.get("brand", "HelloProVision"), footer=style.get("footer", ""),
        accent=accent, accent2=style.get("accent2") or "#F26B5B", accent_dark=_darken(accent), today=date.today().strftime("%Y.%m.%d."),
        version=version, landscape=_wide(content),
    )


def pdf(content: dict, style: dict, version: int = 0) -> bytes:
    from weasyprint import HTML  # a natív könyvtárak betöltése lassú, ezért csak itt

    return HTML(string=html(content, style, version)).write_pdf()


# ── DOCX ─────────────────────────────────────────────────


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def docx(content: dict, style: dict, version: int = 0) -> bytes:
    doc = DocxDocument()
    sec = doc.sections[0]
    if _wide(content):
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    accent = RGBColor.from_string((style.get("accent") or "#A4DA4C").lstrip("#").upper())

    head = doc.add_paragraph()
    r = head.add_run((content.get("chip") or "").upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0x5C, 0x7A, 0x1F)
    doc.add_heading(content.get("title", ""), level=0)
    if content.get("subtitle"):
        p = doc.add_paragraph(content["subtitle"])
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x6B, 0x66)
    meta = doc.add_paragraph(f"{style.get('brand', '')} · {date.today():%Y.%m.%d.}" + (f" · v{version}" if version else ""))
    meta.runs[0].font.size = Pt(8.5)
    if content.get("lead"):
        lead = doc.add_paragraph(content["lead"])
        lead.runs[0].italic = True

    for i, s in enumerate(content.get("sections", []), start=1):
        doc.add_heading(f"{i}. {s['title']}", level=1)
        for b in s.get("blocks", []):
            t = b.get("type")
            if t == "paragraph":
                doc.add_paragraph(b["text"])
            elif t == "bullets":
                for it in b["items"]:
                    doc.add_paragraph(it, style="List Bullet")
            elif t == "callout":
                p = doc.add_paragraph()
                if b.get("title"):
                    rr = p.add_run(b["title"] + ": ")
                    rr.bold = True
                p.add_run(b.get("text", ""))
            elif t in ("kpis",):
                p = doc.add_paragraph()
                for k in b["items"]:
                    rr = p.add_run(f"{k['value']} ")
                    rr.bold, rr.font.size, rr.font.color.rgb = True, Pt(14), accent
                    p.add_run(f"{k['label']}" + (f" ({k['hint']})" if k.get("hint") else "") + "    ")
            elif t in ("cards", "steps"):
                for n, k in enumerate(b["items"], start=1):
                    p = doc.add_paragraph()
                    rr = p.add_run((f"{n}. " if t == "steps" else "") + k["title"] + " – ")
                    rr.bold = True
                    p.add_run(k.get("text", ""))
            elif t == "table":
                cols = b["columns"]
                table = doc.add_table(rows=1, cols=len(cols))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr = table.rows[0]
                _repeat_header(hdr)
                for ci, h in enumerate(cols):
                    cell = hdr.cells[ci]
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(str(h))
                    run.bold, run.font.size = True, Pt(9)
                    run.font.color.rgb = RGBColor(0xF2, 0xF1, 0xEA)
                    _shade(cell, "121210")
                for row in b["rows"]:
                    cells = table.add_row().cells
                    for ci, v in enumerate(row[: len(cols)]):
                        cells[ci].text = ""
                        run = cells[ci].paragraphs[0].add_run("" if v is None else str(v))
                        run.font.size = Pt(9)
                doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── XLSX ─────────────────────────────────────────────────


def xlsx_generic(content: dict) -> bytes:
    """Minden táblás szakasz külön munkalapra; az első lapon a szöveges összefoglaló."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Összefoglaló"
    ws["A1"] = content.get("title", "")
    ws["A1"].font = xlsx.TITLE_FONT
    ws["A2"] = content.get("subtitle", "")
    row = 4
    if content.get("lead"):
        ws.cell(row=row, column=1, value=content["lead"]).alignment = xlsx.WRAP
        row += 2
    ws.column_dimensions["A"].width = 120
    used = {"Összefoglaló"}
    for s in content.get("sections", []):
        texts = [b["text"] for b in s["blocks"] if b.get("type") in ("paragraph", "callout")]
        texts += [i for b in s["blocks"] if b.get("type") == "bullets" for i in b["items"]]
        if texts:
            ws.cell(row=row, column=1, value=s["title"]).font = Font(bold=True, size=12)
            row += 1
            for t in texts:
                ws.cell(row=row, column=1, value=t).alignment = xlsx.WRAP
                row += 1
            row += 1
        for n, b in enumerate([b for b in s["blocks"] if b.get("type") == "table"]):
            name = (s["title"][:28] + (f" {n + 1}" if n else "")).replace("/", "-").replace(":", "")[:31]
            while name in used:
                name = name[:29] + "_" + str(len(used))
            used.add(name)
            t = wb.create_sheet(name)
            xlsx.write_table(t, 1, b["columns"], b["rows"], widths=[min(60, max(12, len(str(c)) + 4)) for c in b["columns"]],
                             wrap_cols=range(1, len(b["columns"]) + 1))
    return xlsx.to_bytes(wb)


CONTENT_STRATEGY_EXPLAIN = [
    ("Hónap", "A tervezett publikálás hónapja."),
    ("Prioritás", "P1 = erős üzleti szándék, közvetlen belső link a szolgáltatási / városi oldalakra; P2 = támogató vagy szélesebb téma."),
    ("Elsődleges kulcsszó", "Az a kulcsszó, amire a cikk épül. A városi kereskedelmi kulcsszavakat a landingek célozzák."),
    ("URL", "A tervezett cikk-URL."),
    ("Javasolt cím", "Munkacím; a szövegíró finomíthatja, de az elsődleges kulcsszó maradjon benne."),
    ("Tartalmi irány", "Mit válaszoljon meg a cikk, és miben legyen jobb, mint a versenytársaké."),
    ("Fő CTA", "A cikk végén ide vezetjük a látogatót."),
]


def content_strategy_xlsx(content: dict, extra: dict) -> bytes:
    """A 6 lapos tartalomstratégia-munkafüzet (a HelloProVision saját táblájának szerkezetében)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Összefoglaló"
    ws["A1"] = content.get("title", "")
    ws["A1"].font = xlsx.TITLE_FONT
    ws["A2"] = content.get("subtitle", "")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    r = 4
    if content.get("lead"):
        ws.cell(row=r, column=1, value="Összefoglaló").font = Font(bold=True)
        ws.cell(row=r, column=2, value=content["lead"]).alignment = xlsx.WRAP
        r += 2
    sections = {s["key"]: s for s in content.get("sections", [])}
    summary = sections.get("summary", {"blocks": []})
    for b in summary["blocks"]:
        if b["type"] == "kpis":
            for k in b["items"]:
                ws.cell(row=r, column=1, value=k["label"]).font = Font(bold=True)
                ws.cell(row=r, column=2, value=k["value"])
                r += 1
            r += 1
        elif b["type"] == "cards":
            for k in b["items"]:
                ws.cell(row=r, column=1, value=k["title"]).font = Font(bold=True)
                ws.cell(row=r, column=2, value=k["text"]).alignment = xlsx.WRAP
                r += 1
            r += 1
        elif b["type"] == "paragraph":
            ws.cell(row=r, column=2, value=b["text"]).alignment = xlsx.WRAP
            r += 1
        elif b["type"] == "table":
            r = xlsx.write_table(ws, r, b["columns"], b["rows"], freeze=False, autofilter=False) + 2
    r += 1
    ws.cell(row=r, column=1, value="Oszlopmagyarázat (6 havi roadmap)").font = Font(bold=True, size=12)
    r += 1
    for col, expl in CONTENT_STRATEGY_EXPLAIN:
        ws.cell(row=r, column=1, value=col).font = Font(bold=True)
        ws.cell(row=r, column=2, value=expl).alignment = xlsx.WRAP
        r += 1

    def table_sheet(title: str, key: str, widths: list[float]):
        s = sections.get(key)
        t = wb.create_sheet(title[:31])
        tables = [b for b in (s or {"blocks": []})["blocks"] if b["type"] == "table"]
        if not tables:
            t["A1"] = "Nincs adat."
            return
        b = tables[0]
        xlsx.write_table(t, 1, b["columns"], b["rows"], widths=widths, wrap_cols=range(1, len(b["columns"]) + 1))

    table_sheet("Kulcsszóklaszterek", "clusters", [34, 10, 8, 18, 22, 50, 14, 36, 10])
    table_sheet(f"{extra.get('months', 6)} havi roadmap", "roadmap", [16, 10, 30, 38, 44, 60, 30])
    table_sheet("Mérési terv", "measurement", [18, 30, 50, 40, 40])
    paid = wb.create_sheet("PPC")
    rows = extra.get("paid") or []
    if rows:
        xlsx.write_table(paid, 1, ["Hónap", "Tartalom", "SEO szerep", "Meta kreatív", "Paid szerep", "Search cél", "Remarketing / következő lépés", "KPI"], rows,
                         widths=[16, 40, 30, 40, 30, 30, 36, 24], wrap_cols=range(1, 9))
    else:
        paid["A1"] = "Nincs paid terv ehhez a stratégiához."
    table_sheet("Következő témák", "parked", [40, 60, 50])
    for sh in wb.worksheets[1:]:
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    return xlsx.to_bytes(wb)


FORMATS: dict[str, dict[str, Any]] = {
    "pdf": {"mime": "application/pdf"},
    "docx": {"mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    "xlsx": {"mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
}
