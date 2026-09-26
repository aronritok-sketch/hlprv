"""Dokumentumgenerálás: adatok → (Claude-szöveg vagy sablonszöveg) → tartalom-JSON → PDF / DOCX / XLSX.

Elv: a számok, URL-ek, kulcsszavak és táblák az adatbázisból jönnek (docfacts), a Claude csak a magyarázó szöveget
írja a HelloProVision hangján, a feltöltött saját minták (referenciadokumentumok) alapján. API-kulcs nélkül is
készül dokumentum: ilyenkor az adatokból összerakott sablonszöveg kerül bele, és a felület ezt jelzi.
"""

import copy
import io
import json
from typing import Optional

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..jobs.runner import JobError, Progress, handler
from ..models import DOC_TYPES, Document, DocumentFile, Job, PaidPlanItem, Project, ReferenceDoc, RoadmapItem, StoredFile
from . import claude, docfacts, render, storage
from . import settings as app_settings
from .activity import log
from .docfacts import num
from .llm import LLMError, claude_available
from .prompts import LANG_NAMES, METHOD_RULES, project_context
from .projects import get_project

PROMPT_VERSION = "doc-2"

FORMATS_BY_TYPE = {
    "seo_strategy": ["pdf", "docx"],
    "content_strategy": ["pdf", "xlsx", "docx"],
    "roadmap": ["pdf", "xlsx"],
    "wireframe_deck": ["pdf", "docx"],
    "dev_brief": ["docx", "pdf"],
    "writer_brief": ["docx", "pdf"],
    "designer_brief": ["docx", "pdf"],
    "seo_checklist": ["xlsx", "pdf"],
    "tech_audit": ["pdf", "docx", "xlsx"],
    "monthly_report": ["pdf", "docx"],
}

# Mi kell az egyes dokumentumokhoz (ha hiányzik, a generálás érthető hibával áll meg).
REQUIRES = {
    "seo_strategy": ("analysis",),
    "content_strategy": ("roadmap",),
    "roadmap": ("roadmap",),
    "wireframe_deck": ("wireframes",),
    "dev_brief": ("structure",),
    "writer_brief": ("structure",),
    "designer_brief": ("structure",),
    "seo_checklist": ("structure",),
    "tech_audit": ("audit",),
}
MISSING = {
    "analysis": "Előbb futtasd le a kulcsszó-elemzést.",
    "structure": "Előbb generáld le az oldalstruktúrát.",
    "roadmap": "Előbb generáld le a tartalmi roadmapet.",
    "wireframes": "Még nincs wireframe; előbb generálj legalább egyet.",
    "audit": "Még nincs technikai audit adat: tölts fel Screaming Frog exportot, vagy futtass crawlt.",
}

HOUSE_STYLE = """A HelloProVision dokumentumstílusa:
- Közvetlen, szakmai, magabiztos hang; többes szám első személy („javasoljuk”, „azt látjuk”). Nincs töltelékszöveg, nincs marketingzsargon.
- Minden állítás mögött adat vagy konkrét megfigyelés áll. Számot csak a megadott adatokból használj; ha nincs adat, ne találj ki.
- Rövid bekezdések (2–4 mondat). Döntést és indoklást írj, ne általánosságokat.
- Nem ígérünk helyezést, forgalmat vagy garantált eredményt; a keresési számok irányadók.
- Az ügyféldokumentum az ügyfélnek szól: szakszavak helyett érthető magyarázat. A belső brief a kollégának szól: konkrét teendők, elfogadási feltételek.
- A táblákat a rendszer illeszti be: ne ismételd meg a tábla tartalmát felsorolásban, csak értelmezd.
"""

NARRATIVE_SCHEMA_BASE = {
    "type": "object",
    "additionalProperties": False,
    "required": ["lead", "sections"],
    "properties": {
        "lead": {"type": "string", "description": "2–3 mondatos bevezető a dokumentum elejére."},
        "sections": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["key", "paragraphs", "bullets"],
                "properties": {
                    "key": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                },
            },
        },
    },
}


def narrative_schema(keys: list[str]) -> dict:
    s = copy.deepcopy(NARRATIVE_SCHEMA_BASE)
    s["properties"]["sections"]["items"]["properties"]["key"] = {"type": "string", "enum": keys}
    return s


HU_ONLY = {"tech_audit"}


def default_language(project: Project, doc_type: str) -> str:
    if DOC_TYPES[doc_type][1] == "internal" or doc_type in HU_ONLY:
        return "hu"
    return "hu" if (project.content_language or "hu").startswith("hu") else "en"


def reference_excerpts(db: Session, doc_type: str, language: str, limit_chars: int = 18000) -> list[tuple[str, str]]:
    refs = db.scalars(
        select(ReferenceDoc).where(ReferenceDoc.is_active.is_(True), ReferenceDoc.doc_type.in_([doc_type, "general"]))
        .order_by((ReferenceDoc.doc_type == doc_type).desc(), (ReferenceDoc.language == language).desc(), ReferenceDoc.id.desc())
    ).all()
    out, used = [], 0
    for r in refs:
        text = (r.extracted_text or "").strip()
        if not text:
            continue
        take = text[: max(0, min(8000, limit_chars - used))]
        if not take:
            break
        out.append((r.title, take))
        used += len(take)
    return out


def system_prompt(db: Session, doc_type: str, language: str) -> str:
    label, audience = DOC_TYPES[doc_type]
    parts = [
        "A HelloProVision (digitális ügynökség, SEO és weboldal-fejlesztés) szakmai dokumentumait írod.",
        METHOD_RULES,
        HOUSE_STYLE,
        f"Dokumentum: {label} ({'ügyfélnek' if audience == 'client' else 'belső, a csapatnak'}).",
    ]
    refs = reference_excerpts(db, doc_type, language)
    if refs:
        parts.append("Saját korábbi dokumentumaink – a hangnemet, a felépítést és a részletességet ezekből vedd át (a tartalmukat ne másold):")
        for title, text in refs:
            parts.append(f"<minta cím=\"{title}\">\n{text}\n</minta>")
    return "\n\n".join(parts)


def user_prompt(project: Project, profile, doc_type: str, language: str, facts: dict, content: dict) -> str:
    sections = [{"key": s["key"], "title": s["title"], "tables": [b["columns"] for b in s["blocks"] if b.get("type") == "table"]}
                for s in content["sections"] if s.get("narrative")]
    return "\n\n".join([
        f"Nyelv: {LANG_NAMES.get(language, language)} – a teljes szöveget ezen a nyelven írd.",
        "Projekt:\n" + project_context(project, profile),
        "Adatok (csak ezekre hivatkozz):\n" + json.dumps(facts, ensure_ascii=False, default=str)[:60000],
        "Szakaszok, amelyekhez szöveget kérünk (a táblákat a rendszer teszi be a szöveg után):\n" + json.dumps(sections, ensure_ascii=False),
        "Szakaszonként 1–3 bekezdés; felsorolás csak ott, ahol valóban lépések vagy szempontok vannak (egyébként üres lista). "
        "A „lead” a teljes dokumentum rövid lényege.",
    ])


# ── Sablonszöveg (ha nincs Claude) ───────────────────────


def template_narrative(doc_type: str, language: str, facts: dict, content: dict) -> dict:
    hu = language == "hu"
    out: dict[str, list[str]] = {}
    lead = ""
    if doc_type == "seo_strategy":
        pr = facts.get("priorities", {})
        cl = facts.get("clusters") or []
        lead = (f"{num(facts.get('keywords', 0))} kulcsszót elemeztünk a(z) {facts.get('domain')} számára. A stratégia a keresési szándékra és az üzleti értékre épül: "
                f"{pr.get('P1', 0)} kulcsszó kapott P1 prioritást, ezekre készülnek a dedikált oldalak és az első hónapok tartalmai."
                if hu else
                f"We analysed {num(facts.get('keywords', 0))} keywords for {facts.get('domain')}. The strategy is built on search intent and business value: "
                f"{pr.get('P1', 0)} keywords are P1 and receive dedicated pages and the first months of content.")
        out["executive_summary"] = [
            ("A javasolt struktúra minden fontos kereskedelmi szándéknak saját URL-t ad, így az oldalak nem versenyeznek egymással. "
             "A tartalmi roadmap ezekre az oldalakra vezeti a látogatókat." if hu else
             "The proposed structure gives each important commercial intent its own URL, so pages do not compete with each other. "
             "The content roadmap drives visitors to these pages.")]
        out["current_situation"] = [
            (f"A kutatott kulcsszavak közül {len(facts.get('own_rankings') or [])} esetben látunk jelenlegi rangsorolást." if facts.get("own_rankings") else
             "A kutatás alapján a domain a vizsgált kulcsszavakra jelenleg nem jelenik meg a találatok elején – ez egyben a legnagyobb lehetőség.")
            if hu else
            (f"The domain currently ranks for {len(facts.get('own_rankings') or [])} of the researched keywords." if facts.get("own_rankings") else
             "The domain does not currently rank near the top for the researched keywords – which is also the main opportunity.")]
        if cl:
            out["market_opportunity"] = [
                (f"A legnagyobb keresleti témakör: {cl[0]['name']} (kb. {num(cl[0]['volume'])} keresés/hó)." if hu else
                 f"The largest demand cluster is {cl[0]['name']} (approx. {num(cl[0]['volume'])} searches/month).")]
        out["keyword_opportunities"] = [
            ("Az alábbi kulcsszavaknál a legerősebb az üzleti szándék; a volumen csak az utolsó szempont." if hu else
             "These keywords carry the strongest business intent; volume is the last criterion, not the first.")]
        out["website_structure"] = [
            ("Egy elsődleges kulcsszó = egy URL. A főoldal a márkát és a fő ajánlatot képviseli, a szolgáltatási és városi oldalak a konkrét kereséseket." if hu else
             "One primary keyword = one URL. The homepage represents the brand, service and city pages target the specific searches.")]
        out["content_roadmap"] = [
            ("A cikkek információs és döntési kérdésekre válaszolnak, majd a releváns szolgáltatási oldalra vezetnek." if hu else
             "Articles answer informational and decision-stage questions, then lead to the relevant service page.")]
    elif doc_type == "content_strategy":
        lead = (f"{facts.get('months')} hónapos tartalmi terv, havonta {facts.get('per_month')} tartalommal: {facts.get('articles')} cikk és {facts.get('case_studies')} case study."
                if hu else
                f"A {facts.get('months')}-month content plan with {facts.get('per_month')} pieces per month: {facts.get('articles')} articles and {facts.get('case_studies')} case studies.")
        out["summary"] = [("A témák a jóváhagyott kulcsszóklaszterekből jönnek; minden cikknek egy elsődleges kulcsszava és egy céloldala van." if hu else
                           "Topics come from the approved keyword clusters; each article has one primary keyword and one target page.")]
        out["roadmap"] = [("A sorrendet a prioritás és a szezonalitás adja; a P1 témák kerülnek előre." if hu else "Order follows priority and seasonality; P1 topics come first.")]
        out["measurement"] = [("Az első hónapokban indexelést és megjelenéseket mérünk, később kattintást és konverziót." if hu else
                               "In the first months we track indexing and impressions, later clicks and conversions.")]
    elif doc_type == "roadmap":
        lead = "A tartalmak havi bontásban." if hu else "Content by month."
    elif doc_type == "wireframe_deck":
        lead = ("Az oldalak felépítése a keresési szándékot és a konverziót szolgálja: minden blokknak van feladata." if hu else
                "Every page layout serves search intent and conversion: each section has a job.")
    elif doc_type == "dev_brief":
        lead = "A jóváhagyott struktúra és wireframe-ek alapján elvégzendő SEO-fejlesztések. Minden sor akkor kész, ha az elfogadási feltétel teljesül."
        out["decision_basis"] = ["A dokumentum a jóváhagyott kulcsszókutatásra, oldalstruktúrára és wireframe-ekre épül. Eltérés esetén előbb egyeztess a SEO managerrel."]
        out["implementation_notes"] = ["Élesítés előtt a staging oldalt futtasd le Screaming Froggal (4xx, átirányítási lánc, canonical, duplikált title/H1), és csak hibamentes állapotban publikálj."]
    elif doc_type == "writer_brief":
        lead = "Oldalankénti kulcsszó-mapping és szövegírói szabályok. Egy elsődleges kulcsszó = egy URL."
        out["principle"] = ["Szolgáltatási oldalnál a szolgáltatáskereső szándék elsőbbséget élvez a volumennel szemben."]
    elif doc_type == "tech_audit":
        from .audit_rules import TOPICS

        xl = [TOPICS[t["topic"]]["title"].lower() for t in facts.get("topics", []) if t["size"] == "XL" and t["findings"]]
        lead = (f"A(z) {facts.get('domain')} technikai állapotát Screaming Froggal és helyszíni ellenőrzésekkel vizsgáltuk. "
                + (f"A legsürgetőbb (XL) teendők: {', '.join(xl)}." if xl else "Kritikus (XL) hibát nem találtunk."))
    elif doc_type == "monthly_report":
        lead = (f"A(z) {facts.get('domain')} havi összefoglalója: {len(facts.get('done', []))} elvégzett feladat, {len(facts.get('published', []))} megjelent tartalom, "
                f"{len(facts.get('fixed', []))} javított technikai hiba." if hu else
                f"Monthly summary for {facts.get('domain')}: {len(facts.get('done', []))} tasks completed, {len(facts.get('published', []))} pieces published, "
                f"{len(facts.get('fixed', []))} technical fixes.")
        out["summary"] = [("A következő hónapban a roadmap szerinti tartalmak és a nyitott technikai feladatok folytatódnak." if hu else
                           "Next month we continue with the roadmap content and the open technical tasks.")]
    elif doc_type == "designer_brief":
        lead = "Oldalcélok, kötelező szakaszok és képi irány a jóváhagyott wireframe-ek alapján."
    return {"lead": lead, "sections": [{"key": k, "paragraphs": v, "bullets": []} for k, v in out.items()]}


def merge_narrative(content: dict, narrative: dict) -> dict:
    c = copy.deepcopy(content)
    c["lead"] = narrative.get("lead", "")
    by_key = {s["key"]: s for s in narrative.get("sections", [])}
    for s in c["sections"]:
        n = by_key.get(s["key"])
        if not n:
            continue
        pre = [{"type": "paragraph", "text": p} for p in n.get("paragraphs", []) if p.strip()]
        if n.get("bullets"):
            pre.append({"type": "bullets", "items": n["bullets"]})
        slot = next((i for i, b in enumerate(s["blocks"]) if b.get("type") == "narrative_slot"), None)
        s["blocks"] = s["blocks"][:slot] + pre + s["blocks"][slot + 1:] if slot is not None else pre + s["blocks"]
    for s in c["sections"]:
        s["blocks"] = [b for b in s["blocks"] if b.get("type") != "narrative_slot"]
    c["sections"] = [s for s in c["sections"] if s["blocks"]]
    return c


# ── Fájlok ───────────────────────────────────────────────


def doc_style(db: Session, language: str) -> dict:
    return {
        "brand": app_settings.get(db, "brand_name") or "HelloProVision",
        "accent": app_settings.get(db, "doc_accent"),
        "accent2": app_settings.get(db, "doc_accent_2"),
        "footer": app_settings.get(db, "doc_footer"),
        "lang": language,
    }


def filename(project: Project, doc: Document, fmt: str) -> str:
    return f"{project.domain} – {DOC_TYPES[doc.doc_type][0]} v{doc.version}.{fmt}"


def render_files(db: Session, project: Project, doc: Document, formats: list[str], user_id: Optional[int]) -> None:
    style = doc_style(db, doc.language)
    old_ids = [f.file_id for f in doc.files]
    doc.files.clear()
    db.flush()
    for fid in old_ids:
        stored = db.get(StoredFile, fid)
        if stored:
            storage.delete(stored)
            db.delete(stored)
    db.flush()
    for fmt in formats:
        if fmt == "pdf":
            data = render.pdf(doc.content, style, doc.version)
        elif fmt == "docx":
            data = render.docx(doc.content, style, doc.version)
        elif fmt == "xlsx":
            if doc.doc_type == "content_strategy":
                paid = db.execute(select(PaidPlanItem).where(PaidPlanItem.project_id == project.id)).scalars().all()
                items = {r.id: r for r in db.scalars(select(RoadmapItem).where(RoadmapItem.project_id == project.id)).all()}
                rows = [[items[p.roadmap_item_id].month.strftime("%Y.%m.") if p.roadmap_item_id in items else "", items[p.roadmap_item_id].title if p.roadmap_item_id in items else "",
                         p.seo_role, p.meta_creative, p.paid_role, p.search_target, p.remarketing_next, p.kpi] for p in paid]
                data = render.content_strategy_xlsx(doc.content, {"months": project.strategy_months, "paid": rows})
            else:
                data = render.xlsx_generic(doc.content)
        else:
            continue
        stored = storage.save(db, data, filename(project, doc, fmt), "document", project.id, user_id, render.FORMATS[fmt]["mime"])
        doc.files.append(DocumentFile(format=fmt, file_id=stored.id))
    db.flush()


def current_hash(db: Session, project: Project, doc_type: str, language: str) -> str:
    builder = docfacts.BUILDERS.get(doc_type)
    if not builder:
        return ""
    facts, _, _ = docfacts.build(db, project, doc_type, language)
    return docfacts.fingerprint(facts)


def check_requirements(d: docfacts.Data, doc_type: str) -> None:
    for need in REQUIRES.get(doc_type, ()):
        if need == "audit":
            from ..models import AuditFinding, AuditTopic

            ok = bool(d.db.scalar(select(func.count()).select_from(AuditFinding).where(AuditFinding.project_id == d.project.id, AuditFinding.count > 0))) or bool(
                d.db.scalar(select(func.count()).select_from(AuditTopic).where(AuditTopic.project_id == d.project.id, AuditTopic.observation != "")))
        else:
            ok = {"analysis": bool(d.analysed), "structure": bool(d.pages), "roadmap": bool(d.roadmap), "wireframes": bool(d.wireframes)}[need]
        if not ok:
            raise JobError(MISSING[need])


@handler("generate_document")
def job_generate_document(db: Session, job: Job, progress: Progress):
    project = get_project(db, job.project_id)
    doc_type = job.payload["doc_type"]
    if doc_type not in docfacts.BUILDERS:
        raise JobError("Ez a dokumentumtípus ebben a modulban nem generálható.")
    language = job.payload.get("language") or default_language(project, doc_type)
    formats = job.payload.get("formats") or FORMATS_BY_TYPE[doc_type]
    progress(0.05, "Adatok összegyűjtése…")
    facts, content, d = docfacts.build(db, project, doc_type, language, job.payload.get("period"))
    check_requirements(d, doc_type)
    method, model, warning = "template", "", ""
    if claude_available(db) and any(s.get("narrative") for s in content["sections"]):
        progress(0.2, "A Claude írja a szöveget…")
        keys = [s["key"] for s in content["sections"] if s.get("narrative")]
        try:
            narrative, model = claude.write_json(
                db, name=f"doc_{doc_type}", version=PROMPT_VERSION, system=system_prompt(db, doc_type, language),
                user=user_prompt(project, d.profile, doc_type, language, facts, content), schema=narrative_schema(keys),
                audience=DOC_TYPES[doc_type][1], project_id=project.id, job_id=job.id,
            )
            method = "claude"
        except LLMError as e:
            warning = str(e)
            narrative = template_narrative(doc_type, language, facts, content)
    else:
        narrative = template_narrative(doc_type, language, facts, content)
    progress(0.7, "Fájlok készítése…")
    version = (db.scalar(select(func.max(Document.version)).where(Document.project_id == project.id, Document.doc_type == doc_type)) or 0) + 1
    doc = Document(
        project_id=project.id, doc_type=doc_type, audience=DOC_TYPES[doc_type][1], language=language, version=version,
        title=content["title"], content=merge_narrative(content, narrative), source_hash=docfacts.fingerprint(facts),
        method=method, model=model, prompt_version=PROMPT_VERSION, created_by=job.created_by,
    )
    db.add(doc)
    db.flush()
    render_files(db, project, doc, formats, job.created_by)
    log(db, project.id, job.created_by, "document", doc.id, "generated", f"{DOC_TYPES[doc_type][0]} v{version} ({method})")
    return {"document_id": doc.id, "version": version, "method": method, "warning": warning}


# ── Referenciadokumentumok ───────────────────────────────


def extract_text(data: bytes, name: str) -> str:
    lower = name.lower()
    if lower.endswith(".docx"):
        from docx import Document as D

        d = D(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if lower.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:  # pragma: no cover
            raise ValueError("A PDF feldolgozásához telepíteni kell a pypdf csomagot.")
        r = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    if lower.endswith((".txt", ".md")):
        for enc in ("utf-8", "cp1250", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    raise ValueError("Támogatott formátum: DOCX, PDF, TXT, MD.")
