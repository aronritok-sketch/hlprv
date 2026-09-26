"""XLSX exportok a HelloProVision saját tábláinak elrendezésében.

Kulcsszókutatás: „Kulcsszókutatás ÉÉÉÉ.HH.” munkalap (kulcsszó × lokáció sorok, versenytársanként Position/URL/Traffic)
és „kulcsszavak lokáció nélkül” lista. Ha már van AI-elemzés, a kulcsszótérkép oszlopai is bekerülnek.
"""

import io
from datetime import date
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1B1B19")
HEADER_FONT = Font(bold=True, color="E8E6D9")
ACCENT_FILL = PatternFill("solid", fgColor="B3B07A")
ACCENT_FONT = Font(bold=True, color="16160F")
TITLE_FONT = Font(bold=True, size=16)
THIN = Side(style="thin", color="D5D3C6")
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, row: int, ncols: int, accent: bool = False) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = ACCENT_FILL if accent else HEADER_FILL
        cell.font = ACCENT_FONT if accent else HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)


def set_widths(ws, widths: dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def write_table(ws, start_row: int, headers: list[str], rows: Iterable[list[Any]], widths: Optional[list[float]] = None,
                wrap_cols: Iterable[int] = (), accent: bool = False, freeze: bool = True, autofilter: bool = True) -> int:
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers), accent)
    r = start_row
    wrap_set = set(wrap_cols)
    for r_i, row in enumerate(rows, start=start_row + 1):
        r = r_i
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in wrap_set:
                cell.alignment = WRAP
    if widths:
        set_widths(ws, {i + 1: w for i, w in enumerate(widths)})
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    if autofilter and r > start_row:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{r}"
    return r


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def keyword_research(project, keywords: list, competitors: list[str], analysis: dict[int, dict] | None = None,
                     include_translation: bool = False) -> bytes:
    """A saját kulcsszókutatás tábla. `keywords`: Keyword ORM objektumok (metrics, rankings betöltve)."""
    from .research import _rank

    analysis = analysis or {}
    own = (project.domain or "").lower()
    wb = Workbook()
    ws = wb.active
    ws.title = f"Kulcsszókutatás {date.today():%Y.%m.}"[:31]

    headers = ["Keyword", "Local long tail keywords", "Parent topic", "Volume", "Keyword Difficulty", "Location"]
    if include_translation:
        headers.insert(1, "Magyar")
    has_analysis = bool(analysis)
    if has_analysis:
        headers += ["Intent", "Cluster", "Priority", "Primary / Secondary", "Recommended URL", "Notes"]
    domains = [own] + [c for c in competitors if c != own]
    for d in domains:
        label = d + "/"
        headers += [f"{label}: Position", f"{label}: URL", f"{label}: Traffic"]

    rows = []
    for kw in keywords:
        by_loc: dict[str, Any] = {}
        for m in sorted(kw.metrics, key=lambda m: _rank(m.source)):
            by_loc.setdefault(m.location, m)
        if not by_loc:
            by_loc = {"": None}
        # Ha van lokációs adat, az országos sor csak akkor jelenik meg, ha nincs más.
        locs = [loc for loc in by_loc if loc] or [""]
        for loc in locs:
            m = by_loc.get(loc) or by_loc.get("")
            row: list[Any] = [kw.term]
            if include_translation:
                row.append(kw.translation)
            vol = m.volume if m and m.volume is not None else _any(kw, "volume")
            kd = m.kd if m and m.kd is not None else _any(kw, "kd")
            row += [kw.modifier, kw.parent_topic, vol, kd, loc]
            if has_analysis:
                a = analysis.get(kw.id, {})
                row += [a.get("intent", ""), a.get("cluster", ""), a.get("priority", ""), a.get("role", ""),
                        a.get("url", ""), a.get("notes", "")]
            for d in domains:
                rk = next((r for r in kw.rankings if r.domain == d and r.location in (loc, "")), None)
                row += [rk.position if rk else None, rk.url if rk else "", rk.traffic if rk and rk.traffic is not None else 0]
            rows.append(row)
    vol_col = headers.index("Volume")
    rows.sort(key=lambda r: (-(r[vol_col] or 0), r[0]))
    widths = [34] + ([30] if include_translation else []) + [18, 22, 10, 12, 16]
    if has_analysis:
        widths += [20, 26, 9, 12, 34, 30]
    widths += [10, 40, 10] * len(domains)
    write_table(ws, 1, headers, rows, widths)

    ws2 = wb.create_sheet("kulcsszavak lokáció nélkül")
    seen, uniq = set(), []
    for r in rows:
        if r[0].lower() not in seen:
            seen.add(r[0].lower())
            uniq.append([r[0]])
    write_table(ws2, 1, ["Keyword"], uniq, [40], freeze=False)
    return to_bytes(wb)


def _any(kw, fld):
    vals = [getattr(m, fld) for m in kw.metrics if getattr(m, fld) is not None]
    return max(vals) if vals and fld == "volume" else (vals[0] if vals else None)
